"""Tests for the UNDP HDR spec-driven bulk-CSV + XLSX fetcher (enrichment stage 2).

Pure logic only — no AWS, no network.  All HTTP calls are replaced with local
fixtures (monkeypatch / mock).  The base's AWS clients are lazy; neither the base
nor fetch_undp requires credentials or internet access under test.

Synthetic payloads:
  - HDR25 wide CSV structure (existing seed codes + new enrichment prefixes)
  - OPHI XLSX Table 6 (MPI time-series)

Existing tests are preserved unchanged; new test classes extend coverage for:
  - INEQ_INC  (ineq_inc_ prefix, new compound-prefix enrichment code)
  - MPI        (OPHI XLSX source, parse_ophi_mpi_xlsx)
  - spec-driven fetch_one routing (enrichment codes via specs.json)
  - apply_derivation (annual_change, binary, unknown → ValueError)
  - fetch_one with mocked HTTP for both HDR CSV and OPHI XLSX paths
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

import numpy as np
import openpyxl
import pandas as pd
import pytest

_QGI_PIPELINE = Path(__file__).resolve().parents[3]
_BASE_DIR = _QGI_PIPELINE / "batch" / "_fetcher_base"
_UNDP_SRC = _QGI_PIPELINE / "batch" / "indicator-fetcher-undp" / "src"
for _p in (_BASE_DIR, _UNDP_SRC):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import qgi_fetcher_base as base  # noqa: E402
import fetch_undp as f  # noqa: E402


# ---------------------------------------------------------------------------
# Shared fixtures
# ---------------------------------------------------------------------------

ISO3 = {
    "AFG": "Afghanistan",
    "NOR": "Norway",
    "BRA": "Brazil",
    "DEU": "Germany",
}


def _make_wide_df() -> pd.DataFrame:
    """Synthetic wide CSV that mirrors the real HDR25 CSV structure.

    Includes all original columns plus new enrichment columns (ineq_inc_).
    No mpi_ columns — matches production HDR25 CSV reality.
    """
    return pd.DataFrame(
        {
            "iso3": ["AFG", "NOR", "BRA", "ZZZ", "DEU"],  # ZZZ = unknown
            "country": ["Afghanistan", "Norway", "Brazil", "Unknown Territory", "Germany"],
            # HDI columns
            "hdi_1990": [0.295, 0.847, 0.603, 0.500, 0.801],
            "hdi_2022": [0.462, 0.966, 0.760, 0.600, 0.950],
            # Gender-disaggregated HDI — should be EXCLUDED from UD-01 parse
            "hdi_f_1990": [0.250, 0.840, 0.580, None, 0.790],
            "hdi_m_1990": [0.330, 0.854, 0.620, None, 0.812],
            # IHDI (only later years have data)
            "ihdi_2010": [0.340, 0.880, 0.570, None, 0.860],
            "ihdi_2022": [0.389, 0.904, 0.583, None, 0.880],
            # GII
            "gii_1995": [0.650, 0.060, 0.407, 0.300, 0.041],
            "gii_2022": [0.655, 0.020, 0.337, 0.250, 0.055],
            # MYS
            "mys_2000": [2.0, 11.8, 6.4, None, 13.2],
            "mys_2022": [3.9, 13.0, 8.1, None, 14.1],
            # EYS
            "eys_2000": [7.1, 16.0, 14.2, None, 15.8],
            "eys_2022": [9.8, 18.1, 15.7, None, 17.0],
            # GDI
            "gdi_2000": [0.620, 0.990, 0.950, None, 0.980],
            "gdi_2022": [0.680, 0.993, 0.972, None, 0.985],
            # INEQ_INC (Atkinson income inequality — compound prefix)
            "ineq_inc_2010": [0.42, 0.07, 0.30, None, 0.09],
            "ineq_inc_2022": [0.45, 0.06, 0.31, None, 0.08],
            # NOTE: no mpi_ columns — matches production HDR25 CSV reality
        }
    )


def _make_ophi_xlsx_bytes() -> bytes:
    """Synthetic OPHI Table 6 XLSX with recognisable column headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table 6"
    # Header row
    ws.append(["iso3", "country_name", "year", "MPI", "headcount_ratio", "intensity"])
    # Data rows  (4 resolvable iso3, 2 years each, 1 unknown)
    rows = [
        ("AFG", "Afghanistan", 2016, 0.272, 0.544, 0.500),
        ("AFG", "Afghanistan", 2022, 0.210, 0.420, 0.500),
        ("NOR", "Norway", 2019, 0.003, 0.006, 0.500),
        ("NOR", "Norway", 2023, 0.002, 0.004, 0.500),
        ("BRA", "Brazil", 2015, 0.017, 0.034, 0.500),
        ("BRA", "Brazil", 2021, 0.012, 0.024, 0.500),
        ("DEU", "Germany", 2018, 0.004, 0.008, 0.500),
        ("DEU", "Germany", 2022, 0.003, 0.006, 0.500),
        ("ZZZ", "Unknown", 2020, 0.100, 0.200, 0.500),  # unknown iso3
    ]
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Tests: parse_undp_wide — HDI (UD-01 prefix "hdi_")  [original tests]
# ---------------------------------------------------------------------------


class TestParseUndpWideHdi:
    def test_basic_melt_and_iso3_resolution(self):
        wide = _make_wide_df()
        out = f.parse_undp_wide(wide, "hdi_", ISO3)
        assert list(out.columns) == ["country", "year", "value"]
        # 4 resolvable ISO3 x 2 years = 8 rows (ZZZ dropped)
        assert len(out) == 8

    def test_values_are_correct(self):
        out = f.parse_undp_wide(_make_wide_df(), "hdi_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Afghanistan", 1990)] == pytest.approx(0.295)
        assert recs[("Norway", 2022)] == pytest.approx(0.966)
        assert recs[("Brazil", 1990)] == pytest.approx(0.603)
        assert recs[("Germany", 2022)] == pytest.approx(0.950)

    def test_gender_disaggregated_columns_excluded(self):
        """hdi_f_* and hdi_m_* must NOT appear as years (their suffix isn't purely numeric)."""
        out = f.parse_undp_wide(_make_wide_df(), "hdi_", ISO3)
        assert set(out["year"].unique()) == {1990, 2022}

    def test_unknown_iso3_dropped(self):
        out = f.parse_undp_wide(_make_wide_df(), "hdi_", ISO3)
        assert "Unknown Territory" not in out["country"].values
        assert all(c in ISO3.values() for c in out["country"].unique())


# ---------------------------------------------------------------------------
# Tests: parse_undp_wide — IHDI (UD-02 prefix "ihdi_")  [original]
# ---------------------------------------------------------------------------


class TestParseUndpWideIhdi:
    def test_ihdi_rows_parsed(self):
        out = f.parse_undp_wide(_make_wide_df(), "ihdi_", ISO3)
        assert len(out) == 8  # 4 countries x 2 years
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Norway", 2022)] == pytest.approx(0.904)
        assert recs[("Afghanistan", 2010)] == pytest.approx(0.340)


# ---------------------------------------------------------------------------
# Tests: parse_undp_wide — GII (UD-03 prefix "gii_")  [original]
# ---------------------------------------------------------------------------


class TestParseUndpWideGii:
    def test_gii_values(self):
        out = f.parse_undp_wide(_make_wide_df(), "gii_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Norway", 1995)] == pytest.approx(0.060)
        assert recs[("Afghanistan", 2022)] == pytest.approx(0.655)


# ---------------------------------------------------------------------------
# Tests: parse_undp_wide — MPI absent (UD-04 scaffolded)  [original]
# ---------------------------------------------------------------------------


class TestParseUndpWideMpi:
    def test_mpi_prefix_absent_returns_empty(self):
        """HDR25 CSV has no mpi_ columns; should return empty frame gracefully."""
        out = f.parse_undp_wide(_make_wide_df(), "mpi_", ISO3)
        assert list(out.columns) == ["country", "year", "value"]
        assert len(out) == 0

    def test_mpi_prefix_present_if_future_csv_added(self):
        """Verify parse logic works if mpi_ columns appear in a future CSV."""
        wide = _make_wide_df().copy()
        wide["mpi_2015"] = [0.30, 0.01, 0.02, 0.40, 0.005]
        wide["mpi_2020"] = [0.28, np.nan, 0.02, 0.38, 0.004]
        out = f.parse_undp_wide(wide, "mpi_", ISO3)
        assert "mpi_2015" not in out.columns  # melted into year/value
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Afghanistan", 2015)] == pytest.approx(0.30)
        # NaN rows dropped (Norway 2020)
        assert ("Norway", 2020) not in recs


# ---------------------------------------------------------------------------
# Tests: parse_undp_wide — MYS (UD-05) and EYS (UD-06)  [original]
# ---------------------------------------------------------------------------


class TestParseUndpWideMysEys:
    def test_mys_values(self):
        out = f.parse_undp_wide(_make_wide_df(), "mys_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Germany", 2022)] == pytest.approx(14.1)
        assert recs[("Afghanistan", 2000)] == pytest.approx(2.0)

    def test_eys_values(self):
        out = f.parse_undp_wide(_make_wide_df(), "eys_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Norway", 2022)] == pytest.approx(18.1)
        assert recs[("Brazil", 2000)] == pytest.approx(14.2)


# ---------------------------------------------------------------------------
# Tests: parse_undp_wide — GDI (UD-07)  [original]
# ---------------------------------------------------------------------------


class TestParseUndpWideGdi:
    def test_gdi_values(self):
        out = f.parse_undp_wide(_make_wide_df(), "gdi_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Norway", 2022)] == pytest.approx(0.993)
        assert recs[("Afghanistan", 2000)] == pytest.approx(0.620)


# ---------------------------------------------------------------------------
# Tests: parse_undp_wide — INEQ_INC (new compound prefix "ineq_inc_")  [NEW]
# ---------------------------------------------------------------------------


class TestParseUndpWideIneqInc:
    """INEQ_INC uses the compound prefix ineq_inc_ which contains an underscore.
    The startswith + 4-digit suffix check must work for multi-part prefixes."""

    def test_ineq_inc_rows_parsed(self):
        out = f.parse_undp_wide(_make_wide_df(), "ineq_inc_", ISO3)
        assert list(out.columns) == ["country", "year", "value"]
        # ZZZ not in ISO3 map; ZZZ row has None -> only 4 countries resolved
        # but ZZZ row in fixture has None value -> 4 countries x 2 years = 8 rows
        assert len(out) == 8

    def test_ineq_inc_values_correct(self):
        out = f.parse_undp_wide(_make_wide_df(), "ineq_inc_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Afghanistan", 2010)] == pytest.approx(0.42)
        assert recs[("Norway", 2022)] == pytest.approx(0.06)
        assert recs[("Germany", 2010)] == pytest.approx(0.09)
        assert recs[("Brazil", 2022)] == pytest.approx(0.31)

    def test_ineq_inc_unknown_iso3_dropped(self):
        out = f.parse_undp_wide(_make_wide_df(), "ineq_inc_", ISO3)
        assert "Unknown Territory" not in out["country"].values

    def test_ineq_inc_years_are_integers(self):
        out = f.parse_undp_wide(_make_wide_df(), "ineq_inc_", ISO3)
        assert set(out["year"].unique()) == {2010, 2022}

    def test_ineq_inc_no_false_columns(self):
        """Columns like ineq_le_ and ineq_edu_ must NOT bleed into ineq_inc_ parse."""
        wide = _make_wide_df().copy()
        wide["ineq_le_2010"] = [0.10, 0.02, 0.05, None, 0.03]
        wide["ineq_edu_2010"] = [0.20, 0.03, 0.08, None, 0.04]
        out = f.parse_undp_wide(wide, "ineq_inc_", ISO3)
        # Should still be 8 rows (2 ineq_inc_ years only, no bleed from ineq_le_ / ineq_edu_)
        assert set(out["year"].unique()) == {2010, 2022}
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        # Value for (Afghanistan, 2010) must be ineq_inc_ value (0.42), not ineq_le_ (0.10)
        assert recs[("Afghanistan", 2010)] == pytest.approx(0.42)


# ---------------------------------------------------------------------------
# Tests: parse_ophi_mpi_xlsx  [NEW]
# ---------------------------------------------------------------------------


class TestParseOphiMpiXlsx:
    def test_basic_parse_returns_expected_rows(self):
        xlsx_bytes = _make_ophi_xlsx_bytes()
        out = f.parse_ophi_mpi_xlsx(xlsx_bytes, ISO3)
        assert list(out.columns) == ["country", "year", "value"]
        # 4 resolvable countries x 2 years = 8 rows; ZZZ dropped
        assert len(out) == 8

    def test_mpi_values_correct(self):
        out = f.parse_ophi_mpi_xlsx(_make_ophi_xlsx_bytes(), ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Afghanistan", 2016)] == pytest.approx(0.272)
        assert recs[("Afghanistan", 2022)] == pytest.approx(0.210)
        assert recs[("Norway", 2019)] == pytest.approx(0.003)
        assert recs[("Germany", 2018)] == pytest.approx(0.004)

    def test_unknown_iso3_dropped(self):
        out = f.parse_ophi_mpi_xlsx(_make_ophi_xlsx_bytes(), ISO3)
        assert "Unknown" not in out["country"].values
        assert all(c in ISO3.values() for c in out["country"].unique())

    def test_year_is_integer(self):
        out = f.parse_ophi_mpi_xlsx(_make_ophi_xlsx_bytes(), ISO3)
        assert out["year"].dtype == np.int64 or out["year"].dtype == object or set(map(type, out["year"])) == {int}
        # All years must be int-like
        for y in out["year"]:
            assert int(y) == y

    def test_empty_xlsx_returns_empty(self):
        """Corrupt/empty XLSX bytes should return empty frame, not crash."""
        out = f.parse_ophi_mpi_xlsx(b"not an xlsx", ISO3)
        assert list(out.columns) == ["country", "year", "value"]
        assert len(out) == 0

    def test_deduplication_averages(self):
        """Duplicate (iso3, year) rows should be averaged."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["iso3", "country_name", "year", "MPI"])
        ws.append(["NOR", "Norway", 2020, 0.002])
        ws.append(["NOR", "Norway", 2020, 0.004])  # duplicate year
        buf = io.BytesIO()
        wb.save(buf)
        out = f.parse_ophi_mpi_xlsx(buf.getvalue(), ISO3)
        assert len(out) == 1
        assert out.iloc[0]["value"] == pytest.approx(0.003)

    def test_nan_mpi_values_dropped(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["iso3", "country_name", "year", "MPI"])
        ws.append(["AFG", "Afghanistan", 2020, None])
        ws.append(["NOR", "Norway", 2020, 0.003])
        buf = io.BytesIO()
        wb.save(buf)
        out = f.parse_ophi_mpi_xlsx(buf.getvalue(), ISO3)
        assert len(out) == 1
        assert out.iloc[0]["country"] == "Norway"


# ---------------------------------------------------------------------------
# Tests: apply_derivation  [NEW]
# ---------------------------------------------------------------------------


class TestApplyDerivation:
    def _base_df(self) -> pd.DataFrame:
        return pd.DataFrame(
            {
                "country": ["Norway", "Norway", "Norway", "Afghanistan", "Afghanistan"],
                "year": [2000, 2001, 2002, 2000, 2001],
                "value": [100.0, 110.0, 121.0, 50.0, 55.0],
            }
        )

    def test_none_derivation_returns_unchanged(self):
        df = self._base_df()
        out = f.apply_derivation(df, "none")
        assert len(out) == len(df)
        pd.testing.assert_frame_equal(out, df)

    def test_null_string_derivation_returns_unchanged(self):
        df = self._base_df()
        out = f.apply_derivation(df, "null")
        assert len(out) == len(df)

    def test_empty_string_derivation_returns_unchanged(self):
        df = self._base_df()
        out = f.apply_derivation(df, "")
        assert len(out) == len(df)

    def test_annual_change_correct_diffs(self):
        df = self._base_df()
        out = f.apply_derivation(df, "annual_change")
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        # First year per country is dropped (NaN diff)
        assert ("Norway", 2000) not in recs
        assert ("Afghanistan", 2000) not in recs
        assert recs[("Norway", 2001)] == pytest.approx(10.0)
        assert recs[("Norway", 2002)] == pytest.approx(11.0)
        assert recs[("Afghanistan", 2001)] == pytest.approx(5.0)

    def test_annual_change_drops_nan_first_row(self):
        df = self._base_df()
        out = f.apply_derivation(df, "annual_change")
        # 5 rows → 3 after dropping first-of-each-country
        assert len(out) == 3

    def test_binary_threshold(self):
        df = pd.DataFrame(
            {
                "country": ["A", "A", "A"],
                "year": [2000, 2001, 2002],
                "value": [0.3, 0.5, 0.7],
            }
        )
        out = f.apply_derivation(df, "binary:0.5")
        vals = list(out["value"])
        assert vals == pytest.approx([0.0, 1.0, 1.0])

    def test_binary_threshold_float(self):
        df = pd.DataFrame(
            {
                "country": ["X"],
                "year": [2000],
                "value": [0.499999],
            }
        )
        out = f.apply_derivation(df, "binary:0.5")
        assert out.iloc[0]["value"] == pytest.approx(0.0)

    def test_unsupported_derivation_raises(self):
        df = self._base_df()
        with pytest.raises(ValueError, match="unsupported derivation"):
            f.apply_derivation(df, "ratio:HDI:IHDI")

    def test_unsupported_per_1000_pop_raises(self):
        df = self._base_df()
        with pytest.raises(ValueError, match="unsupported derivation"):
            f.apply_derivation(df, "per_1000_pop")

    def test_binary_bad_threshold_raises(self):
        df = self._base_df()
        with pytest.raises(ValueError):
            f.apply_derivation(df, "binary:not_a_number")


# ---------------------------------------------------------------------------
# Tests: spec-driven fetch_one routing (mocked HTTP + iso3_map)  [NEW]
# ---------------------------------------------------------------------------


class TestFetchOneEnrichmentCodes:
    """Integration tests for the enrichment code path in fetch_one.

    HTTP is replaced by monkeypatching fetch_undp._load_wide_df and
    fetch_undp._load_xlsx_bytes.  load_iso3_country_map is replaced by a
    mock returning the shared ISO3 fixture.
    """

    def setup_method(self):
        """Clear module-level caches before each test to avoid cross-test leakage."""
        f._wide_df_cache.clear()
        f._xlsx_cache.clear()
        f._enrichment_specs = None  # force reload from specs.json

    def _patch_csv(self, monkeypatch):
        monkeypatch.setattr(f, "_load_wide_df", lambda: _make_wide_df())

    def _patch_iso3(self, monkeypatch):
        monkeypatch.setattr(base, "load_iso3_country_map", lambda: ISO3)

    def test_enrichment_IHDI_routes_to_csv(self, monkeypatch):
        self._patch_csv(monkeypatch)
        self._patch_iso3(monkeypatch)
        out = f.fetch_one({"id_string": "IHDI"})
        assert out is not None
        assert list(out.columns) == ["country", "year", "value"]
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Norway", 2022)] == pytest.approx(0.904)

    def test_enrichment_GII_routes_to_csv(self, monkeypatch):
        self._patch_csv(monkeypatch)
        self._patch_iso3(monkeypatch)
        out = f.fetch_one({"id_string": "GII"})
        assert out is not None
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Afghanistan", 2022)] == pytest.approx(0.655)

    def test_enrichment_INEQ_INC_routes_to_csv(self, monkeypatch):
        """NEW code INEQ_INC — compound prefix ineq_inc_; fetchable from HDR25 CSV."""
        self._patch_csv(monkeypatch)
        self._patch_iso3(monkeypatch)
        out = f.fetch_one({"id_string": "INEQ_INC"})
        assert out is not None
        assert list(out.columns) == ["country", "year", "value"]
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Afghanistan", 2010)] == pytest.approx(0.42)
        assert recs[("Germany", 2022)] == pytest.approx(0.08)

    def test_enrichment_MPI_routes_to_xlsx(self, monkeypatch):
        """NEW code MPI — routed to ophi_xlsx source; parse_ophi_mpi_xlsx called."""
        xlsx_bytes = _make_ophi_xlsx_bytes()
        monkeypatch.setattr(f, "_load_xlsx_bytes", lambda url: xlsx_bytes)
        self._patch_iso3(monkeypatch)
        out = f.fetch_one({"id_string": "MPI"})
        assert out is not None
        assert list(out.columns) == ["country", "year", "value"]
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Afghanistan", 2016)] == pytest.approx(0.272)
        assert recs[("Norway", 2023)] == pytest.approx(0.002)

    def test_seed_UD01_still_works(self, monkeypatch):
        """Seed code UD-01 must still route through INDICATOR_PREFIX."""
        self._patch_csv(monkeypatch)
        self._patch_iso3(monkeypatch)
        out = f.fetch_one({"id_string": "UD-01"})
        assert out is not None
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert recs[("Norway", 2022)] == pytest.approx(0.966)

    def test_seed_UD04_mpi_returns_none(self, monkeypatch):
        """UD-04 (scaffolded mpi_ in seed set) returns None when prefix absent."""
        self._patch_csv(monkeypatch)
        self._patch_iso3(monkeypatch)
        out = f.fetch_one({"id_string": "UD-04"})
        assert out is None

    def test_unknown_code_raises(self, monkeypatch):
        self._patch_csv(monkeypatch)
        self._patch_iso3(monkeypatch)
        with pytest.raises(ValueError, match="unknown UNDP indicator"):
            f.fetch_one({"id_string": "ZZ-99"})


# ---------------------------------------------------------------------------
# Tests: specs.json loaded correctly  [NEW]
# ---------------------------------------------------------------------------


class TestSpecsJson:
    def test_specs_file_exists(self):
        assert (_UNDP_SRC / "specs.json").exists(), "specs.json missing from src/"

    def test_all_enrichment_codes_present(self):
        specs = f._load_enrichment_specs()
        for code in ("IHDI", "GII", "INEQ_INC", "MPI"):
            assert code in specs, f"specs.json missing enrichment code {code!r}"

    def test_all_fetchable(self):
        specs = f._load_enrichment_specs()
        for code, entry in specs.items():
            assert entry.get("fetchable") is True, f"{code} is not fetchable but present in specs"

    def test_hdr_csv_specs_have_col_prefix(self):
        specs = f._load_enrichment_specs()
        for code, entry in specs.items():
            if entry.get("source") == "hdr_composite_csv":
                assert entry.get("col_prefix"), f"{code} hdr_composite_csv spec missing col_prefix"

    def test_mpi_source_is_ophi_xlsx(self):
        specs = f._load_enrichment_specs()
        assert specs["MPI"]["source"] == "ophi_xlsx"

    def test_ineq_inc_col_prefix_correct(self):
        specs = f._load_enrichment_specs()
        assert specs["INEQ_INC"]["col_prefix"] == "ineq_inc_"


# ---------------------------------------------------------------------------
# Tests: edge cases  [original + extended]
# ---------------------------------------------------------------------------


class TestEdgeCases:
    def test_nan_values_dropped(self):
        wide = _make_wide_df().copy()
        wide.loc[wide["iso3"] == "NOR", "hdi_1990"] = np.nan
        out = f.parse_undp_wide(wide, "hdi_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert ("Norway", 1990) not in recs
        assert ("Norway", 2022) in recs  # other year still present

    def test_non_numeric_value_coerced_to_nan_and_dropped(self):
        wide = _make_wide_df().copy()
        wide["hdi_1990"] = wide["hdi_1990"].astype(object)
        wide.loc[wide["iso3"] == "BRA", "hdi_1990"] = "n.a."
        out = f.parse_undp_wide(wide, "hdi_", ISO3)
        recs = {(r["country"], r["year"]): r["value"] for r in out.to_dict("records")}
        assert ("Brazil", 1990) not in recs

    def test_empty_wide_df_returns_empty(self):
        empty = pd.DataFrame(columns=_make_wide_df().columns)
        out = f.parse_undp_wide(empty, "hdi_", ISO3)
        assert list(out.columns) == ["country", "year", "value"]
        assert len(out) == 0

    def test_all_unknown_iso3_returns_empty(self):
        wide = _make_wide_df().copy()
        wide["iso3"] = ["ZZZ", "ZZA", "ZZB", "ZZC", "ZZD"]
        out = f.parse_undp_wide(wide, "hdi_", ISO3)
        assert len(out) == 0

    def test_deduplication_via_mean(self):
        wide = pd.DataFrame(
            {
                "iso3": ["NOR", "NOR"],
                "country": ["Norway", "Norway"],
                "hdi_2000": [0.80, 0.90],
            }
        )
        out = f.parse_undp_wide(wide, "hdi_", ISO3)
        assert len(out) == 1
        assert out.iloc[0]["value"] == pytest.approx(0.85)

    def test_output_columns_exactly(self):
        out = f.parse_undp_wide(_make_wide_df(), "hdi_", ISO3)
        assert list(out.columns) == ["country", "year", "value"]


# ---------------------------------------------------------------------------
# Tests: INDICATOR_PREFIX catalog covers all 7 seed specs  [original]
# ---------------------------------------------------------------------------


class TestIndicatorCatalog:
    def test_all_seven_indicators_wired(self):
        expected = {"UD-01", "UD-02", "UD-03", "UD-04", "UD-05", "UD-06", "UD-07"}
        assert set(f.INDICATOR_PREFIX) == expected

    def test_known_prefixes(self):
        assert f.INDICATOR_PREFIX["UD-01"] == "hdi_"
        assert f.INDICATOR_PREFIX["UD-02"] == "ihdi_"
        assert f.INDICATOR_PREFIX["UD-03"] == "gii_"
        assert f.INDICATOR_PREFIX["UD-04"] == "mpi_"
        assert f.INDICATOR_PREFIX["UD-05"] == "mys_"
        assert f.INDICATOR_PREFIX["UD-06"] == "eys_"
        assert f.INDICATOR_PREFIX["UD-07"] == "gdi_"


# ---------------------------------------------------------------------------
# Tests: shared base functions  [original]
# ---------------------------------------------------------------------------


class TestSharedBase:
    def test_iso3_map_from_df_uppercases_and_skips_blanks(self):
        df = pd.DataFrame(
            {
                "iso_alpha3": ["afg", "USA", "", None],
                "qgi_name": ["Afghanistan", "United States", "Nowhere", "AlsoNowhere"],
            }
        )
        m = base.iso3_map_from_df(df)
        assert m == {"AFG": "Afghanistan", "USA": "United States"}

    def test_to_canonical_frame_coerces_dedups_sorts_drops_nan(self):
        df = pd.DataFrame(
            {
                "country": ["B", "A", "A", "A"],
                "year": ["2001", 2000, 2000, 2002],
                "value": [1.0, 5.0, 9.0, np.nan],
            }
        )
        out = base.to_canonical_frame(df)
        assert list(out.columns) == base.OUTPUT_COLUMNS
        assert out["year"].dtype == np.int64
        assert out["value"].dtype == np.float64
        assert (out["is_predicted"] == False).all()  # noqa: E712
        assert out.to_dict("records") == [
            {"country": "A", "year": 2000, "value": 9.0, "is_predicted": False},
            {"country": "B", "year": 2001, "value": 1.0, "is_predicted": False},
        ]

    def test_to_canonical_frame_requires_core_columns(self):
        with pytest.raises(ValueError):
            base.to_canonical_frame(pd.DataFrame({"country": ["A"], "year": [2000]}))


def test_fetch_one_rejects_unknown_id():
    with pytest.raises(ValueError):
        f.fetch_one({"id_string": "ZZ-99"})
